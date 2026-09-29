#!/usr/bin/env python3
#
# Copyright (c) 2026, Centre Tecnologic de Telecomunicacions de Catalunya (CTTC)
#
# SPDX-License-Identifier: GPL-2.0-only
#
# Author: Gabriel Ferreira <gabrielcarvfer@gmail.com>
"""Generate the reference CDFs of three-gpp-channel-calibration-test-suite.cc.

3GPP reference: the full calibration of TR 38.901 Sec. 7.8.2 (Table 7.8-2, BS
antenna configuration 1), whose results (based on TR 38.900 V14.0.0) are the
per-company CDFs of R1-165975, "E-mail discussion summary of the full scale
calibration", RAN1#85, attachment Phase2Config1Calibration_v35_CMCC.xlsx:

    https://www.3gpp.org/ftp/tsg_ran/WG1_RL1/TSGR1_85/Docs/R1-165975.zip

Each sheet (one per scenario and carrier frequency) holds, for every metric, one
column per company with the 0..100 percentiles in rows 29..129. For every
percentile, the header stores the mean across companies and the band spanned by
them (minimum and maximum).

Sionna reference (optional): the delay and angle spreads of the serving links of
the TR 38.901 channel models of Sionna (https://github.com/NVlabs/sionna), on the
calibration drops of sionna.sys, with the serving site attached on the path loss
without shadow fading, the angle spread of TR 38.901 Annex A and the parameters
of TR 38.901 V16.1.0 (equal to V15.0.0 for these scenarios). For every
percentile, the header stores the Sionna percentile and the band of its
confidence interval (three standard deviations of the order statistic).

Usage:
    three-gpp-channel-calibration-reference.py R1-165975.zip [--sionna] > three-gpp-channel-calibration-reference.h
"""

import argparse
import io
import math
import statistics
import sys
import zipfile

import openpyxl

SHEET = "Phase2Config1Calibration_v35_CMCC.xlsx"
SCENARIOS = ["UMa", "UMi", "InH"]
FREQUENCIES = [6, 30, 60, 70]
# Metric label prefixes of row 28 of the sheets, and the C++ enumerators.
METRICS = [
    ("Coupling loss", "COUPLING_LOSS"),
    ("Wideband SIR", "SIR"),
    ("Delay Spread", "DS"),
    ("ASD", "ASD"),
    ("ZSD", "ZSD"),
    ("ASA", "ASA"),
    ("ZSA", "ZSA"),
]
PERCENTILES = list(range(5, 100, 5))
FIRST_PERCENTILE_ROW = 29
HEADER_ROW = 28
SOURCE_ROW = 25

# Sionna drops: scenarios and carrier frequencies of the test, and the number of
# drops per building type (low and high loss, pooled into 50%/50%).
SIONNA_FREQUENCIES = [6, 30]
SIONNA_DROPS = 4


def metric_columns(ws):
    """Map each metric to the columns of its per-company CDFs."""
    columns = {}
    label = None
    for col in range(1, ws.max_column + 1):
        head = ws.cell(HEADER_ROW, col).value
        if head not in (None, " "):
            label = str(head)
        source = ws.cell(SOURCE_ROW, col).value
        if label is None or source in (None, 0, "Source", "Mean"):
            continue
        for prefix, name in METRICS:
            if label.startswith(prefix):
                columns.setdefault(name, []).append(col)
    return columns


def percentile_values(ws, col, percentile):
    value = ws.cell(FIRST_PERCENTILE_ROW + percentile, col).value
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_3gpp(path):
    """Mean and band across companies of the R1-165975 CDFs."""
    with zipfile.ZipFile(path) as archive:
        name = next(n for n in archive.namelist() if n.endswith(SHEET))
        workbook = openpyxl.load_workbook(io.BytesIO(archive.read(name)), data_only=True)
    rows = []
    for scenario in SCENARIOS:
        for fc in FREQUENCIES:
            ws = workbook[f"{scenario}-{fc}GHz"]
            columns = metric_columns(ws)
            for _, metric in METRICS:
                mean, low, high = [], [], []
                companies = 0
                for p in PERCENTILES:
                    values = [
                        v
                        for v in (percentile_values(ws, c, p) for c in columns.get(metric, []))
                        if v is not None
                    ]
                    companies = max(companies, len(values))
                    mean.append(statistics.fmean(values))
                    low.append(min(values))
                    high.append(max(values))
                rows.append((scenario, fc, metric, companies, mean, low, high))
    return rows


def sionna_spreads(scenario, fc_ghz, o2i_model, drops):
    """Delay (ns) and angle (deg) spreads of the serving links of Sionna drops."""
    import numpy as np
    import torch
    from sionna.phy.channel.tr38901 import InH, PanelArray, UMa, UMi
    from sionna.phy.channel.tr38901.metrics import (
        angular_spreads_from_rays,
        delay_spread_from_rays,
    )
    from sionna.sys import gen_tr38901_indoor_office_topology, gen_tr38901_multicell_topology

    fc = fc_ghz * 1e9

    def array():
        return PanelArray(
            num_rows_per_panel=1,
            num_cols_per_panel=1,
            polarization="single",
            polarization_type="V",
            antenna_pattern="omni",
            carrier_frequency=fc,
        )

    if scenario == "InH":
        channel = InH(
            carrier_frequency=fc,
            ut_array=array(),
            bs_array=array(),
            direction="downlink",
            indoor_scenario="open",
            spec_version="16.1",
        )
        topology = gen_tr38901_indoor_office_topology(batch_size=drops, num_ut_per_sector=10)
    else:
        model = UMa if scenario == "UMa" else UMi
        channel = model(
            carrier_frequency=fc,
            o2i_model=o2i_model,
            ut_array=array(),
            bs_array=array(),
            direction="downlink",
            spec_version="16.1",
        )
        topology = gen_tr38901_multicell_topology(
            scenario.lower(), batch_size=drops, num_ut_per_sector=10, carrier_frequency=fc
        )
    channel.set_topology(*topology)
    lsp = channel.sample_lsp()
    rays = channel._ray_sampler(lsp)
    serving = torch.argmax(-lsp.pathloss, dim=1)
    out = {"DS": delay_spread_from_rays(rays, lsp, channel._scenario, serving=serving) * 1e9}
    spreads = angular_spreads_from_rays(rays, lsp, channel._scenario, serving=serving)
    for key in ("asd", "zsd", "asa", "zsa"):
        out[key.upper()] = torch.rad2deg(spreads[key])
    return {k: np.asarray(v.flatten().cpu()) for k, v in out.items()}


def read_sionna():
    """Percentiles and their confidence bands of the Sionna spreads."""
    import numpy as np
    from sionna.phy import config

    config.seed = 1
    rows = []
    for scenario in SCENARIOS:
        for fc in SIONNA_FREQUENCIES:
            if scenario == "InH":
                runs = [sionna_spreads(scenario, fc, None, 2 * SIONNA_DROPS)]
            else:
                runs = [sionna_spreads(scenario, fc, m, SIONNA_DROPS) for m in ("low", "high")]
            for metric in ("DS", "ASD", "ZSD", "ASA", "ZSA"):
                x = np.sort(np.concatenate([r[metric] for r in runs]))
                n = len(x)
                value, low, high = [], [], []
                for p in PERCENTILES:
                    q = p / 100
                    d = 3 * math.sqrt(q * (1 - q) / n)
                    value.append(float(np.quantile(x, q)))
                    low.append(float(np.quantile(x, max(q - d, 0))))
                    high.append(float(np.quantile(x, min(q + d, 1))))
                rows.append((scenario, fc, metric, n, value, low, high))
    return rows


def write_table(out, name, rows):
    out.write(f"const std::array<ReferenceCdf, {len(rows)}> {name}{{{{\n")

    def fmt(v):
        return ", ".join(f"{x:.4g}" for x in v)

    for scenario, fc, metric, sources, mean, low, high in rows:
        out.write(f'    {{"{scenario}", {fc}, Metric::{metric}, {sources},\n')
        out.write(f"     {{{fmt(mean)}}},\n     {{{fmt(low)}}},\n     {{{fmt(high)}}}}},\n")
    out.write("}};\n\n")


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("zip", help="path to R1-165975.zip")
    parser.add_argument("--sionna", action="store_true", help="also generate the Sionna reference")
    args = parser.parse_args()

    rows_3gpp = read_3gpp(args.zip)
    rows_sionna = read_sionna() if args.sionna else []

    out = sys.stdout
    out.write(
        "/*\n * Copyright (c) 2026, Centre Tecnologic de Telecomunicacions de Catalunya (CTTC)\n"
    )
    out.write(" *\n * SPDX-License-Identifier: GPL-2.0-only\n */\n\n")
    out.write("// Generated by three-gpp-channel-calibration-reference.py; do not edit.\n\n")
    out.write("#ifndef THREE_GPP_CHANNEL_CALIBRATION_REFERENCE_H\n")
    out.write("#define THREE_GPP_CHANNEL_CALIBRATION_REFERENCE_H\n\n")
    out.write("#include <array>\n#include <cstdint>\n\n")
    out.write("namespace ns3\n{\nnamespace calibration\n{\n\n")
    out.write("// clang-format off\n")
    out.write("/// Calibration metrics of TR 38.901 Table 7.8-2 (configuration 1)\n")
    out.write("enum class Metric : uint8_t\n{\n")
    for _, name in METRICS:
        out.write(f"    {name},\n")
    out.write("};\n\n")
    n = len(PERCENTILES)
    out.write("/// Reference CDF of a calibration metric, at kReferencePercentiles\n")
    out.write("struct ReferenceCdf\n{\n")
    out.write("    const char* scenario;         ///< scenario: UMa, UMi or InH\n")
    out.write("    double fcGHz;                 ///< carrier frequency in GHz\n")
    out.write("    Metric metric;                ///< the metric\n")
    out.write("    uint32_t sources;             ///< number of companies, or of samples\n")
    out.write(f"    std::array<double, {n}> mean; ///< reference value per percentile\n")
    out.write(f"    std::array<double, {n}> low;  ///< lower bound of the reference band\n")
    out.write(f"    std::array<double, {n}> high; ///< upper bound of the reference band\n")
    out.write("};\n\n")
    out.write(
        f"/// Percentiles of the reference CDFs\nconstexpr std::array<double, {n}> "
        f"kReferencePercentiles{{{', '.join(str(p) for p in PERCENTILES)}}};\n\n"
    )
    out.write(
        "/**\n * TR 38.901 Sec. 7.8.2 full calibration (configuration 1) results of R1-165975:\n"
    )
    out.write(" * per percentile, the mean across the companies and the band they span.\n */\n")
    write_table(out, "k3gppFullCalibration", rows_3gpp)
    out.write(
        "/**\n * Delay and angle spreads of the Sionna TR 38.901 channel models, serving site\n"
    )
    out.write(" * attached on the path loss: per percentile, the Sionna value and the band of\n")
    out.write(" * its confidence interval; the number of sources is the number of samples.\n */\n")
    write_table(out, "kSionnaSpreads", rows_sionna)
    out.write("// clang-format on\n\n} // namespace calibration\n} // namespace ns3\n\n")
    out.write("#endif // THREE_GPP_CHANNEL_CALIBRATION_REFERENCE_H\n")


if __name__ == "__main__":
    main()
